"""Shared helpers for the CDS 2.0 deck scripts.

The source deck is assets/deck/CDS-2.0-Creator-Deck-All-Slides.pptx (46 slides:
1-30 standard, 31 divider, 32-46 extras). Never edit it in place; scripts
always write a new file.
"""
from pathlib import Path
from pptx import Presentation

SKILL_DIR = Path(__file__).resolve().parent.parent
SOURCE_DECK = SKILL_DIR / "assets" / "deck" / "CDS-2.0-Creator-Deck-All-Slides.pptx"
EMU_PER_PT = 12700
DIVIDER_SLIDE = 31
CLOSING_SLIDES = {30, 46}
EXTRA_SLIDES = set(range(32, 47))


def walk(shapes):
    """Yield every shape, descending into groups."""
    for sh in shapes:
        if sh.shape_type == 6:  # group
            yield from walk(sh.shapes)
        else:
            yield sh


def font_size(shape):
    sizes = [r.font.size.pt for p in shape.text_frame.paragraphs for r in p.runs if r.font.size]
    return max(sizes) if sizes else 12.0


def capacity(shape):
    """Rough max characters a text box holds at its current font size.
    Zoho Puvi averages ~0.52 em per character; lines are ~1.18 em tall."""
    fs = font_size(shape)
    w = (shape.width or 0) / EMU_PER_PT
    h = (shape.height or 0) / EMU_PER_PT
    per_line = max(1, int(w / (0.52 * fs)))
    lines = max(1, int(h / (1.18 * fs)))
    return int(per_line * lines * 0.9)


def is_furniture(shape):
    """Page numbers and the footer URL are not content slots."""
    t = shape.text_frame.text.strip()
    if t.isdigit() and len(t) <= 2:
        return True
    return t.lower() in {"zoho.com/creator"}


def text_slots(slide):
    out = []
    for sh in walk(slide.shapes):
        if sh.has_text_frame and sh.text_frame.text.strip() and not is_furniture(sh):
            out.append(sh)
    return out


def pictures(slide):
    return [sh for sh in walk(slide.shapes) if sh.shape_type == 13]


def sample_images(slide):
    """Only images the marketer must replace. They are named 'Photo (replace)';
    logos, patterns and overlays are system artwork and must stay."""
    return [p for p in pictures(slide) if "replace" in p.name.lower()]


# Slide 8's three step cards hold drawn sample screens but aren't named "(replace)".
SCREEN_SHAPES_BY_SLIDE = {8: {"Shape 3", "Shape 19", "Shape 41"}}
FIXED_LABELS = {"thank you"}


def screen_regions(slide_no, slide):
    regions = []
    for sh in walk(slide.shapes):
        named = "replace" in sh.name.lower() and sh.shape_type != 13
        if named or sh.name in SCREEN_SHAPES_BY_SLIDE.get(slide_no, set()):
            regions.append((sh.left, sh.top, sh.left + sh.width, sh.top + sh.height))
    return regions


def slot_role(slide_no, slide, shape, regions=None):
    """'screen' = text drawn inside a sample product screen (replace the whole screen),
    'label'  = fixed small caps label or legal line (leave as is),
    'content' = text the marketer's story fills."""
    regions = screen_regions(slide_no, slide) if regions is None else regions
    cx, cy = shape.left + shape.width / 2, shape.top + shape.height / 2
    if any(l <= cx <= r and t <= cy <= b for l, t, r, b in regions):
        return "screen"
    t = shape.text_frame.text.strip()
    if t.startswith("©") or t.lower() in FIXED_LABELS:
        return "label"
    if t.isupper() and len(t) <= 16 and font_size(shape) <= 8:
        return "label"
    return "content"
