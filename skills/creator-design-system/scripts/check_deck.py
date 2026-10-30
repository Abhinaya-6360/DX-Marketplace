"""Brand check for a deck built with build_deck.py.

Usage:
  python check_deck.py plan.json

Reports, in plain language:
  ERRORS   - must be fixed before handing over (text over capacity, structure)
  TO CHECK - things the marketer must replace or confirm (samples, images, numbers)
Exit code 1 if there are errors.
"""
import json
import sys
from pptx import Presentation
from deck_common import (SOURCE_DECK, walk, text_slots, sample_images, capacity, screen_regions, slot_role,
                         DIVIDER_SLIDE, CLOSING_SLIDES)

COVERS = {1, 32, 33}
NUMBER_SLIDES = {5: "stat and its source", 20: "three numbers and their sources",
                 24: "result numbers", 42: "four outcome numbers and their sources",
                 43: "funnel counts", 21: "bar values (bars are shapes: resize by hand)",
                 41: "line values (lines are shapes: redraw by hand)", 22: "chart data",
                 27: "current prices", 19: "current certifications"}
NAMES = {1: "Cover", 16: "AI (Zia)", 30: "Thank you", 31: "divider", 46: "Questions"}


def main(plan_path):
    plan = json.load(open(plan_path))
    prs = Presentation(SOURCE_DECK)
    errors, checks = [], []
    order = [s["source"] for s in plan["slides"]]

    if DIVIDER_SLIDE in order:
        errors.append("Slide 31 (the divider) is still in the deck. Remove it.")
    if order and order[0] not in COVERS:
        errors.append(f"The deck should open with a cover (1, 32 or 33), not slide {order[0]}.")
    if order and order[-1] not in CLOSING_SLIDES:
        errors.append(f"The deck should close with 30 (Thank you) or 46 (Questions), not slide {order[-1]}.")
    if len([n for n in order if n in COVERS]) > 1:
        errors.append("More than one cover slide.")
    if 16 in order:
        checks.append("Slide 16 (AI) is included: keep it only if this deck is about AI (violet is for AI and partner content).")

    for pos, item in enumerate(plan["slides"], 1):
        n = item["source"]
        slide = prs.slides[n - 1]
        by_id = {sh.shape_id: sh for sh in walk(slide.shapes)}
        new = {int(k): v for k, v in item.get("text", {}).items()}
        for sid, text in new.items():
            sh = by_id[sid]
            cap = capacity(sh)
            if len(text.replace("\n", " ")) > cap:
                errors.append(f"Slide {pos} (source {n}): text in slot {sid} is {len(text)} characters; it holds about {cap}. Shorten it — don't shrink the font.")
        regions = screen_regions(n, slide)
        roles = {sh.shape_id: slot_role(n, slide, sh, regions) for sh in text_slots(slide)}
        if any(r == "screen" for r in roles.values()):
            checks.append(f"Slide {pos} (source {n}): the product screen is a drawn sample. Replace it with a real screenshot, or keep it only if it fits the story.")
        left = [sh.text_frame.text.strip() for sh in text_slots(slide)
                if sh.shape_id not in new and roles[sh.shape_id] == "content"]
        if left:
            short = "; ".join(t.replace("\n", " / ")[:40] for t in left[:4])
            checks.append(f"Slide {pos} (source {n}): {len(left)} slot(s) still have sample text — {short}{' …' if len(left) > 4 else ''}")
        imgs = sample_images(slide)
        if imgs:
            checks.append(f"Slide {pos} (source {n}): replace {len(imgs)} sample photo(s) with Replace image.")
        if n in NUMBER_SLIDES:
            checks.append(f"Slide {pos} (source {n}): confirm the {NUMBER_SLIDES[n]}. Never publish a number without a source.")

    print("BRAND CHECK")
    print("Errors:" if errors else "Errors: none")
    for e in errors:
        print("  ✗", e)
    print("To check before sending:" if checks else "To check: none")
    for c in checks:
        print("  •", c)
    sys.exit(1 if errors else 0)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
