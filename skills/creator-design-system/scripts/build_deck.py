"""Build a marketer's deck from the CDS 2.0 source deck.

Usage:
  python build_deck.py plan.json output.pptx

plan.json:
{
  "slides": [
    {"source": 1,  "text": {"12": "Approvals, automated", "15": "Product launch · Nov 2026"},
                   "talk": "Open with the customer's world, not ours."},
    {"source": 5,  "text": {"7": "Approvals take 9 days.", "8": "Your IT team isn't the bottleneck."}},
    {"source": 30}
  ]
}

- "source" is the slide number in the 46-slide deck. Slides appear in plan order;
  every slide not listed is removed (including slide 31, the divider).
- "text" maps shape id (from deck_inventory.py) to new text. Use "\\n" for a new
  paragraph. Formatting comes from the first run of the matching paragraph, so
  fonts, sizes and colours stay on-brand. Shapes not listed keep their sample text.
- "talk" is added to the speaker notes under the slide's own note.
"""
import copy
import json
import sys
from pptx import Presentation
from deck_common import SOURCE_DECK, walk


def set_text(shape, new_text):
    tf = shape.text_frame
    paras = tf.paragraphs
    lines = new_text.split("\n")
    # keep one template paragraph per new line (reuse the last one when there are more lines)
    templates = [copy.deepcopy(p._p) for p in paras]
    txBody = tf._txBody
    for p in list(paras):
        txBody.remove(p._p)
    for i, line in enumerate(lines):
        p_el = copy.deepcopy(templates[min(i, len(templates) - 1)])
        txBody.append(p_el)
    for p, line in zip(tf.paragraphs, lines):
        runs = p.runs
        if runs:
            runs[0].text = line
            for r in runs[1:]:
                r._r.getparent().remove(r._r)
        else:
            p.text = line


def main(plan_path, out_path):
    plan = json.load(open(plan_path))
    prs = Presentation(SOURCE_DECK)
    slides = list(prs.slides)
    wanted = [s["source"] for s in plan["slides"]]
    if len(set(wanted)) != len(wanted):
        sys.exit("Each source slide can appear once. Duplicate slides in Zoho Show after import.")

    for item in plan["slides"]:
        slide = slides[item["source"] - 1]
        by_id = {sh.shape_id: sh for sh in walk(slide.shapes)}
        for sid, text in item.get("text", {}).items():
            sh = by_id.get(int(sid))
            if sh is None or not sh.has_text_frame:
                sys.exit(f"Slide {item['source']}: no text shape with id {sid}")
            set_text(sh, text)
        if item.get("talk"):
            nf = slide.notes_slide.notes_text_frame
            nf.text = (nf.text + "\n\nTalk track: " if nf.text else "Talk track: ") + item["talk"]

    # reorder and drop
    sldIdLst = prs.slides._sldIdLst
    ids = list(sldIdLst)
    keep = [ids[n - 1] for n in wanted]
    for el in ids:
        if el not in keep:
            prs.part.drop_rel(el.rId)
        sldIdLst.remove(el)
    for el in keep:
        sldIdLst.append(el)
    prs.save(out_path)
    print(f"Saved {out_path}: {len(keep)} slides from source {wanted}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
