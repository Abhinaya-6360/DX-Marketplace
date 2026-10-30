"""List the text slots on chosen slides of the CDS 2.0 deck.

Usage:
  python deck_inventory.py 1 5 20 30      # specific slides
  python deck_inventory.py all

Prints JSON: for each slide, every text slot with its shape id, sample text,
font size, max characters and role (content / label / screen), plus picture count and the slide's own note.
Use the shape ids in the build plan (see build_deck.py).
"""
import json
import sys
from pptx import Presentation
from deck_common import SOURCE_DECK, text_slots, sample_images, font_size, capacity, screen_regions, slot_role


def main(args):
    prs = Presentation(SOURCE_DECK)
    nums = range(1, len(prs.slides) + 1) if args == ["all"] else [int(a) for a in args]
    out = []
    for n in nums:
        s = prs.slides[n - 1]
        regions = screen_regions(n, s)
        out.append({
            "slide": n,
            "slots": [{"id": sh.shape_id, "sample": sh.text_frame.text,
                       "size_pt": round(font_size(sh)), "max_chars": capacity(sh),
                       "role": slot_role(n, s, sh, regions)}
                      for sh in text_slots(s)],
            "sample_images": len(sample_images(s)),
            "note": s.notes_slide.notes_text_frame.text if s.has_notes_slide else "",
        })
    print(json.dumps(out, indent=1, ensure_ascii=False))


if __name__ == "__main__":
    main(sys.argv[1:] or ["all"])
